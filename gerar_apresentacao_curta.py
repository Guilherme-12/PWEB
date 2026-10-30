from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUT = r"C:\Users\Guilh\OneDrive\Desktop\PWEB\FitLifeManager.apresentacao-curta.pptx"

GREEN = RGBColor(0x46, 0x7F, 0x1F)
GREEN_DARK = RGBColor(0x35, 0x63, 0x1A)
BLUE = RGBColor(0x1F, 0x3A, 0x2D)
GRAY = RGBColor(0x5B, 0x65, 0x5D)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def text_box(slide, left, top, width, height, paragraphs, size, bold=False, color=BLUE, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    for i, (text, sz, b, col, al) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(sz)
        p.font.bold = b
        p.font.color.rgb = col
        p.font.name = "Calibri"
        p.alignment = al
    return shape


def slide(title, subtitle, bullets, footer):
    slide = prs.slides.add_slide(blank)
    text_box(slide, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.2), [(title, 46, True, GREEN_DARK, PP_ALIGN.LEFT)], 46, True, GREEN_DARK)
    if subtitle:
        text_box(slide, Inches(0.8), Inches(2.7), Inches(11.7), Inches(0.8), [(subtitle, 26, False, GRAY, PP_ALIGN.LEFT)], 26, False, GRAY)
    if bullets:
        text_box(slide, Inches(0.8), Inches(3.6), Inches(11.7), Inches(2.8), bullets, 24, False, BLUE)
    text_box(slide, Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.4), [(footer, 16, False, GRAY, PP_ALIGN.LEFT)], 16, False, GRAY)
    return slide


# Slide 1 - Capa
slide(
    title="FitLife Manager",
    subtitle="Plataforma Web de Gestão do FitLife Gym",
    bullets=[],
    footer="FitLife Gym • Programação Web",
)

# Slide 2 - Contexto / Objetivo
slide(
    title="Contexto e Objetivo",
    subtitle="Do problema à solução",
    bullets=[
        ("Gestão de sócios, planos, aulas, presenças e pagamentos", 24, False, BLUE, PP_ALIGN.LEFT),
        ("Substituir processos manuais por uma plataforma integrada", 24, False, BLUE, PP_ALIGN.LEFT),
        ("Disponibilizar indicadores de gestão", 24, False, BLUE, PP_ALIGN.LEFT),
    ],
    footer="FitLife Manager",
)

# Slide 3 - Funcionalidades
slide(
    title="Funcionalidades",
    subtitle="O que a aplicação faz",
    bullets=[
        ("Dashboard: indicadores e gráficos", 24, False, BLUE, PP_ALIGN.LEFT),
        ("Sócios, planos, aulas, presenças e pagamentos", 24, False, BLUE, PP_ALIGN.LEFT),
        ("Validações de negócio e controlo de capacidade", 24, False, BLUE, PP_ALIGN.LEFT),
    ],
    footer="FitLife Manager",
)

# Slide 4 - Conclusão
slide(
    title="Conclusão",
    subtitle="Projeto concluído",
    bullets=[
        ("Plataforma web funcional", 24, False, BLUE, PP_ALIGN.LEFT),
        ("Dashboard com estatísticas", 24, False, BLUE, PP_ALIGN.LEFT),
        ("Próximos passos: login, relatórios, estados", 24, False, BLUE, PP_ALIGN.LEFT),
    ],
    footer="FitLife Manager",
)

prs.save(OUT)
print("Guardado:", OUT)
print("Total de slides:", len(prs.slides))
